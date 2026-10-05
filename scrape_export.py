"""
Excel export for Instagram scrape results — matches the Flipkart IG Posts workbook layout.
"""

from __future__ import annotations

import io
import json
from datetime import date, datetime

HEADERS = [
    "#",
    "Date (IST)",
    "Time (IST)",
    "Posted By",
    "Paid Partnership",
    "Pinned",
    "Content Type",
    "Caption (Full)",
    "Hashtags",
    "Duration (Sec)",
    "Likes",
    "Comments",
    "Views",
    "Plays",
    "Post Link",
    "Cover Image",
]

COLUMN_WIDTHS = {
    "A": 5.0,
    "B": 13.0,
    "C": 10.0,
    "D": 26.0,
    "E": 11.0,
    "F": 9.0,
    "G": 28.0,
    "H": 60.0,
    "I": 34.0,
    "J": 10.0,
    "K": 11.0,
    "L": 11.0,
    "M": 13.0,
    "N": 14.0,
    "O": 12.0,
    "P": 12.0,
}


def _parse_meta(article):
    raw = article.get("scrape_meta")
    if isinstance(raw, dict):
        return raw
    if raw:
        try:
            return json.loads(raw)
        except (TypeError, json.JSONDecodeError):
            pass
    return {}


def _title_brand(articles, handle=None, label=None):
    handle = (handle or "").strip().lstrip("@")
    if not handle and articles:
        meta = _parse_meta(articles[0])
        handle = (meta.get("profile_handle") or "").strip().lstrip("@")
        label = label or meta.get("profile_label") or ""
    if not handle:
        handle = "instagram"
    brand = (label or "").strip() or f"{handle.title()} Instagram"
    if "instagram" not in brand.lower():
        brand = f"{brand} Instagram"
    brand = brand.replace(" India Instagram", " Instagram")
    dates = []
    for article in articles:
        meta = _parse_meta(article)
        d = meta.get("date_ist") or article.get("pub_date")
        if d:
            dates.append(str(d)[:10])
    if dates:
        dates.sort()
        start = _format_title_date(dates[0])
        end = _format_title_date(dates[-1])
        range_part = f"Posts from {start} to {end}"
    else:
        range_part = f"Posts export {date.today().strftime('%d %b %Y')}"
    return f"{brand} (@{handle}) - {range_part}"


def _format_title_date(iso_date):
    try:
        dt = datetime.strptime(iso_date[:10], "%Y-%m-%d")
        return f"{dt.day} {dt.strftime('%b %Y')}"
    except ValueError:
        return iso_date


def _excel_date(iso_date):
    if not iso_date:
        return None
    try:
        return datetime.strptime(str(iso_date)[:10], "%Y-%m-%d")
    except ValueError:
        return None


def _row_from_article(index, article):
    meta = _parse_meta(article)
    caption = article.get("body") or article.get("title") or ""
    post_url = meta.get("post_url") or article.get("resolved_url") or article.get("url") or ""
    cover = meta.get("cover_image") or article.get("image_url") or ""
    duration = meta.get("duration_sec", "")
    if duration == "" or duration is None:
        duration_val = ""
    else:
        try:
            duration_val = float(duration)
        except (TypeError, ValueError):
            duration_val = duration

    def _num(val):
        if val in ("", None):
            return ""
        try:
            return int(val)
        except (TypeError, ValueError):
            try:
                return float(val)
            except (TypeError, ValueError):
                return val

    return {
        "num": index,
        "date": _excel_date(meta.get("date_ist") or article.get("pub_date")),
        "time": meta.get("time_ist") or "",
        "posted_by": meta.get("posted_by") or "",
        "paid": meta.get("paid_partnership") or "No",
        "pinned": meta.get("pinned") or "No",
        "content_type": meta.get("content_type") or article.get("subtitle") or "",
        "caption": caption,
        "hashtags": meta.get("hashtags") or "",
        "duration": duration_val,
        "likes": _num(meta.get("likes")),
        "comments": _num(meta.get("comments")),
        "views": _num(meta.get("views")),
        "plays": _num(meta.get("plays")),
        "post_url": post_url,
        "cover": cover,
    }


def _sort_key(article):
    meta = _parse_meta(article)
    d = meta.get("date_ist") or article.get("pub_date") or ""
    t = meta.get("time_ist") or "00:00"
    return (str(d), t)


def _article_handle(article):
    meta = _parse_meta(article)
    handle = (meta.get("profile_handle") or "").strip().lstrip("@")
    if handle:
        return handle.lower()
    page = (article.get("page") or "").strip().lstrip("@")
    if page:
        return page.lower()
    source = (article.get("source") or "").strip()
    if source:
        return source.lstrip("@").split()[0].lower()
    return "instagram"


def _write_sheet(wb, articles, *, handle=None, label=None, sheet_title=None, first=False):
    from openpyxl.styles import Font

    articles = list(articles)
    articles.sort(key=_sort_key)
    safe_handle = (handle or "instagram").strip().lstrip("@") or "instagram"
    title = (sheet_title or f"{safe_handle.title()} IG Posts")[:31]
    if first:
        ws = wb.active
        ws.title = title
    else:
        # Avoid duplicate sheet names
        base = title
        n = 2
        existing = {s.title for s in wb.worksheets}
        while title in existing:
            title = f"{base[:28]} {n}"[:31]
            n += 1
        ws = wb.create_sheet(title)

    title_text = _title_brand(articles, handle=handle, label=label)
    ws.merge_cells("A1:P1")
    title_cell = ws["A1"]
    title_cell.value = title_text
    title_cell.font = Font(bold=True, size=12)

    ws.append([])
    ws.append(HEADERS)
    for cell in ws[3]:
        cell.font = Font(bold=True)

    for idx, article in enumerate(articles, start=1):
        row = _row_from_article(idx, article)
        ws.append(
            [
                row["num"],
                row["date"],
                row["time"],
                row["posted_by"],
                row["paid"],
                row["pinned"],
                row["content_type"],
                row["caption"],
                row["hashtags"],
                row["duration"],
                row["likes"],
                row["comments"],
                row["views"],
                row["plays"],
                "Open post",
                "View image",
            ]
        )
        excel_row = ws.max_row
        if row["post_url"]:
            link_cell = ws.cell(excel_row, 15)
            link_cell.hyperlink = row["post_url"]
            link_cell.style = "Hyperlink"
        if row["cover"]:
            img_cell = ws.cell(excel_row, 16)
            img_cell.hyperlink = row["cover"]
            img_cell.style = "Hyperlink"
        date_cell = ws.cell(excel_row, 2)
        if row["date"]:
            date_cell.number_format = "dd-mmm-yyyy"

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width
    return ws


def instagram_articles_to_xlsx_bytes(articles, *, handle=None, label=None, sheet_title=None):
    """Build Flipkart-style workbook. One sheet per handle when exporting all brands."""
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is not installed. Add it to requirements and redeploy."
        ) from exc

    articles = list(articles)
    wb = Workbook()

    if handle:
        _write_sheet(
            wb,
            articles,
            handle=handle,
            label=label,
            sheet_title=sheet_title,
            first=True,
        )
    else:
        groups = {}
        for article in articles:
            key = _article_handle(article)
            groups.setdefault(key, []).append(article)
        if not groups:
            _write_sheet(wb, [], handle="instagram", sheet_title="Instagram IG Posts", first=True)
        else:
            first = True
            for h in sorted(groups.keys()):
                group = groups[h]
                meta = _parse_meta(group[0])
                _write_sheet(
                    wb,
                    group,
                    handle=h,
                    label=meta.get("profile_label") or label,
                    sheet_title=f"{h.title()} IG Posts",
                    first=first,
                )
                first = False

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_filename(handle=None, start_date=None, end_date=None):
    handle = (handle or "Instagram").strip().lstrip("@") or "Instagram"
    brand = handle.title()
    if start_date and end_date:
        try:
            s_dt = datetime.strptime(start_date[:10], "%Y-%m-%d")
            e_dt = datetime.strptime(end_date[:10], "%Y-%m-%d")
            s = f"{s_dt.day}{s_dt.strftime('%b')}".lower()
            e = f"{e_dt.day}{e_dt.strftime('%b%Y')}".lower()
            return f"{brand}_Instagram_{s}-{e}.xlsx"
        except ValueError:
            pass
    return f"{brand}_Instagram_{date.today().isoformat()}.xlsx"
