"""
Export Intelligence Hub reports to HTML and Word (.docx).
"""

from __future__ import annotations

import io
from html import escape as html_escape

from sector_keywords import SECTOR_LABELS


def _section_rows(title, rows, fields):
    lines = [f"<h2>{html_escape(title)}</h2>"]
    if not rows:
        lines.append("<p><em>None identified.</em></p>")
        return "\n".join(lines)
    for row in rows:
        if isinstance(row, str):
            lines.append(f"<p>{html_escape(row)}</p>")
            continue
        headline_parts = [str(row.get(f) or "").strip() for f in fields]
        headline = " — ".join(p for p in headline_parts if p) or "Update"
        lines.append(f"<h3>{html_escape(headline)}</h3>")
        summary = (row.get("summary") or "").strip()
        if summary:
            lines.append(f"<p>{html_escape(summary)}</p>")
    return "\n".join(lines)


def report_to_html(report_row):
    payload = report_row.get("report") or {}
    label = report_row.get("sector_label") or SECTOR_LABELS.get(
        report_row.get("sector"), report_row.get("sector")
    )
    period = report_row.get("period_label") or (
        f"{report_row.get('start_date')} – {report_row.get('end_date')}"
    )
    exec_summary = html_escape((payload.get("executive_summary") or "").strip() or "No summary.")
    sections = [
        _section_rows("Key Metrics", payload.get("key_metrics"), ["company", "metric", "value"]),
        _section_rows("Strategic Moves", payload.get("strategic_moves"), ["company", "move", "importance"]),
        _section_rows("Funding", payload.get("funding"), ["company", "amount", "investor"]),
        _section_rows("Partnerships", payload.get("partnerships"), ["company_a", "company_b", "purpose"]),
        _section_rows("Product Launches", payload.get("product_launches"), ["company", "product"]),
        _section_rows("Competitive Landscape", payload.get("competitive_landscape"), ["company", "development"]),
        _section_rows("What to Watch", payload.get("what_to_watch"), ["watch"]),
    ]
    sources = report_row.get("sources") or []
    source_html = ""
    if sources:
        items = []
        for src in sources:
            title = html_escape(src.get("title") or "Source")
            href = src.get("resolved_url") or src.get("url") or ""
            if href:
                items.append(f'<li><a href="{html_escape(href)}">{title}</a></li>')
            else:
                items.append(f"<li>{title}</li>")
        source_html = "<h2>Source Articles</h2><ul>" + "".join(items) + "</ul>"

    return f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<title>Intelligence — {html_escape(str(label))}</title>
<style>
body {{ font-family: Georgia, 'Times New Roman', serif; max-width: 820px; margin: 40px auto;
  color: #0f172a; line-height: 1.55; padding: 0 20px; }}
h1,h2,h3 {{ font-family: 'Segoe UI', Arial, sans-serif; color: #1e3a5f; }}
.meta {{ color: #64748b; margin-bottom: 24px; }}
.summary {{ background: #f8fafc; border-left: 4px solid #0d9488; padding: 16px 18px; }}
@media print {{ .no-print {{ display: none; }} }}
</style></head><body>
<button class="no-print" onclick="window.print()">Print / Save as PDF</button>
<h1>Insights Hour — {html_escape(str(label))}</h1>
<p class="meta">{html_escape(str(period))}</p>
<h2>Executive Summary</h2>
<div class="summary">{exec_summary}</div>
{''.join(sections)}
{source_html}
</body></html>"""


def report_to_docx_bytes(report_row):
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is not installed. Add it to requirements and redeploy."
        ) from exc

    payload = report_row.get("report") or {}
    label = report_row.get("sector_label") or SECTOR_LABELS.get(
        report_row.get("sector"), report_row.get("sector")
    )
    period = report_row.get("period_label") or (
        f"{report_row.get('start_date')} – {report_row.get('end_date')}"
    )

    doc = Document()
    doc.add_heading(f"Insights Hour — {label}", level=1)
    doc.add_paragraph(str(period))
    doc.add_heading("Executive Summary", level=2)
    doc.add_paragraph((payload.get("executive_summary") or "").strip() or "No summary.")

    def add_section(title, rows, fields):
        doc.add_heading(title, level=2)
        if not rows:
            doc.add_paragraph("None identified.")
            return
        for row in rows:
            if isinstance(row, str):
                doc.add_paragraph(row, style="List Bullet")
                continue
            parts = [str(row.get(f) or "").strip() for f in fields]
            headline = " — ".join(p for p in parts if p) or "Update"
            doc.add_heading(headline, level=3)
            summary = (row.get("summary") or "").strip()
            if summary:
                doc.add_paragraph(summary)

    add_section("Key Metrics", payload.get("key_metrics"), ["company", "metric", "value"])
    add_section("Strategic Moves", payload.get("strategic_moves"), ["company", "move", "importance"])
    add_section("Funding", payload.get("funding"), ["company", "amount", "investor"])
    add_section("Partnerships", payload.get("partnerships"), ["company_a", "company_b", "purpose"])
    add_section("Product Launches", payload.get("product_launches"), ["company", "product"])
    add_section("Competitive Landscape", payload.get("competitive_landscape"), ["company", "development"])
    add_section("What to Watch", payload.get("what_to_watch"), ["watch"])

    sources = report_row.get("sources") or []
    if sources:
        doc.add_heading("Source Articles", level=2)
        for src in sources:
            title = src.get("title") or "Source"
            href = src.get("resolved_url") or src.get("url") or ""
            doc.add_paragraph(f"{title} — {href}" if href else title, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
